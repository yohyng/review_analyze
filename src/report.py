"""Step 9: PowerPoint report generation (python-pptx, native charts/tables).

We use python-pptx *native* charts and tables (not embedded images) so:
  * Japanese text renders correctly when opened in PowerPoint
  * charts/tables stay editable in the deck
  * no kaleido / matplotlib / font dependencies needed

build_report(conn, target_name, axis, insights, output_path) -> Path
"""
from __future__ import annotations

import io
import sqlite3
from datetime import date
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from . import analysis, config, text_analysis, topic_score
from .llm import InsightResult
from .topics import Topic

# --- palette --------------------------------------------------------------- #
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLUE = RGBColor(0x29, 0x80, 0xB9)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GREEN = RGBColor(0x27, 0xAE, 0x60)
RED = RGBColor(0xC0, 0x39, 0x2B)
LIGHT = RGBColor(0xEC, 0xF0, 0xF1)
GREY = RGBColor(0x7F, 0x8C, 0x8D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MAGENTA = RGBColor(0xB0, 0x33, 0x8A)   # VoiceBAUM accent
POS = RGBColor(0x4F, 0x8A, 0x6B)
NEU = RGBColor(0xC9, 0xC3, 0xB6)
NEG = RGBColor(0xC6, 0x6B, 0x61)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# --------------------------------------------------------------------------- #
# low-level slide helpers
# --------------------------------------------------------------------------- #
def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide, left, top, width, height, text, *, size=18, bold=False,
             color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _title_bar(slide, title: str, subtitle: str = ""):
    """Coloured header band + title text, used on every content slide."""
    band = slide.shapes.add_shape(
        1, 0, 0, SLIDE_W, Inches(1.0)  # 1 = MSO_SHAPE.RECTANGLE
    )
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    band.shadow.inherit = False
    tf = band.text_frame
    tf.margin_left = Inches(0.4)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = WHITE
    if subtitle:
        _textbox(slide, Inches(0.4), Inches(1.05), Inches(12.5), Inches(0.4),
                 subtitle, size=13, color=GREY)


def _bullets(slide, left, top, width, height, items, *, size=16,
             marker="・", color=NAVY, space_after=8):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = f"{marker}{item}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def _table(slide, left, top, width, height, headers, rows,
           *, header_fill=NAVY, accent=None):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = shape.table

    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = str(h)
        para = cell.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        para.runs[0].font.size = Pt(13)
        para.runs[0].font.bold = True
        para.runs[0].font.color.rgb = WHITE
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(val)
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(12)
            para.runs[0].font.color.rgb = NAVY
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
            if accent and c == accent["col"]:
                para.runs[0].font.color.rgb = accent["color"]
                para.runs[0].font.bold = True
    return table


# --------------------------------------------------------------------------- #
# content slides
# --------------------------------------------------------------------------- #
def _slide_title(prs, target_name, baseline_label):
    slide = _blank(prs)
    bg = slide.shapes.add_shape(1, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    bg.shadow.inherit = False

    _textbox(slide, Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.2),
             "口コミ分析レポート", size=44, bold=True, color=WHITE)
    _textbox(slide, Inches(0.85), Inches(3.8), Inches(11.7), Inches(0.8),
             f"対象施設：{target_name}", size=24, color=LIGHT)
    _textbox(slide, Inches(0.85), Inches(4.5), Inches(11.7), Inches(0.6),
             f"比較基準：{baseline_label}", size=16, color=GREY)
    _textbox(slide, Inches(0.85), Inches(6.5), Inches(11.7), Inches(0.5),
             f"作成日：{date.today():%Y年%m月%d日}", size=14, color=GREY)


def _slide_summary(prs, target_name, insights, comp):
    slide = _blank(prs)
    _title_bar(slide, "エグゼクティブサマリー")

    _textbox(slide, Inches(0.5), Inches(1.3), Inches(12.3), Inches(1.6),
             insights.summary or "（サマリー未生成）", size=16, color=NAVY)

    # key metric cards
    if comp is not None and not comp.diff.empty:
        top_strength = comp.diff.idxmax()
        top_weak = comp.diff.idxmin()
        cards = [
            ("口コミ件数", f"{insights_extra_reviews(insights)}", BLUE),
            ("最大の強み", f"{top_strength}\n{comp.diff.max():+.1f}pt", GREEN),
            ("最大の弱み", f"{top_weak}\n{comp.diff.min():+.1f}pt", RED),
        ]
    else:
        cards = [("口コミ件数", f"{insights_extra_reviews(insights)}", BLUE)]

    x = Inches(0.5)
    for label, value, color in cards:
        card = slide.shapes.add_shape(1, x, Inches(3.2), Inches(3.9), Inches(2.0))
        card.fill.solid()
        card.fill.fore_color.rgb = LIGHT
        card.line.color.rgb = color
        card.line.width = Pt(2)
        card.shadow.inherit = False
        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        r.font.size = Pt(14); r.font.color.rgb = GREY; r.font.bold = True
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = value
        r2.font.size = Pt(22); r2.font.color.rgb = color; r2.font.bold = True
        x += Inches(4.15)


def insights_extra_reviews(insights) -> str:
    return getattr(insights, "_n_reviews", "-")


def _slide_score_charts(prs, comp):
    slide = _blank(prs)
    _title_bar(slide, "スコア比較", f"{comp.target_label}  vs  {comp.baseline_label}")

    cats = comp.metrics
    # radar (left)
    radar_data = CategoryChartData()
    radar_data.categories = cats
    radar_data.add_series(comp.target_label, [round(comp.target[m], 1) for m in cats])
    radar_data.add_series(comp.baseline_label, [round(comp.baseline[m], 1) for m in cats])
    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.RADAR, Inches(0.4), Inches(1.4), Inches(6.2), Inches(5.4), radar_data
    )
    _style_chart(gf.chart)

    # clustered bar (right)
    bar_data = CategoryChartData()
    bar_data.categories = cats
    bar_data.add_series(comp.target_label, [round(comp.target[m], 1) for m in cats])
    bar_data.add_series(comp.baseline_label, [round(comp.baseline[m], 1) for m in cats])
    gf2 = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.8), Inches(1.4), Inches(6.1), Inches(5.4), bar_data
    )
    _style_chart(gf2.chart)


def _style_chart(chart):
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    try:
        chart.series[0].format.fill.solid()
        chart.series[0].format.fill.fore_color.rgb = BLUE
        chart.series[1].format.fill.solid()
        chart.series[1].format.fill.fore_color.rgb = ORANGE
    except (IndexError, AttributeError):
        pass


def _slide_top5(prs, comp):
    slide = _blank(prs)
    _title_bar(slide, "強み・弱み TOP5", f"比較基準：{comp.baseline_label}")

    strengths, weaknesses = analysis.top_n(comp.diff, n=5)

    _textbox(slide, Inches(0.5), Inches(1.3), Inches(6), Inches(0.4),
             "💪 強み TOP5", size=18, bold=True, color=GREEN)
    s_rows = [[r["指標"], f'{r["差（対象−比較）"]:+.1f}pt'] for _, r in strengths.iterrows()]
    if s_rows:
        _table(slide, Inches(0.5), Inches(1.8), Inches(6), Inches(0.4 * len(s_rows)),
               ["指標", "差"], s_rows, header_fill=GREEN,
               accent={"col": 1, "color": GREEN})

    _textbox(slide, Inches(7.0), Inches(1.3), Inches(6), Inches(0.4),
             "⚠️ 弱み TOP5", size=18, bold=True, color=RED)
    w_rows = [[r["指標"], f'{r["差（対象−比較）"]:+.1f}pt'] for _, r in weaknesses.iterrows()]
    if w_rows:
        _table(slide, Inches(7.0), Inches(1.8), Inches(6), Inches(0.4 * len(w_rows)),
               ["指標", "差"], w_rows, header_fill=RED,
               accent={"col": 1, "color": RED})


def _slide_text_analysis(prs, profile):
    slide = _blank(prs)
    _title_bar(slide, "テキスト分析", f"対象口コミ {profile.n_reviews} 件 / TF-IDF・N-gram")

    _textbox(slide, Inches(0.5), Inches(1.3), Inches(6), Inches(0.4),
             "特徴キーワード（TF-IDF TOP10）", size=16, bold=True, color=BLUE)
    kw = profile.tfidf_keywords.head(10)
    if not kw.empty:
        rows = [[w, f"{s:.3f}"] for w, s in zip(kw["単語"], kw["スコア"])]
        _table(slide, Inches(0.5), Inches(1.8), Inches(6), Inches(4.2),
               ["単語", "スコア"], rows, header_fill=BLUE)

    _textbox(slide, Inches(7.0), Inches(1.3), Inches(6), Inches(0.4),
             "頻出フレーズ（バイグラム TOP10）", size=16, bold=True, color=GREEN)
    bi = profile.bigrams.head(10)
    if not bi.empty:
        rows = [[p, str(c)] for p, c in zip(bi["フレーズ"], bi["件数"])]
        _table(slide, Inches(7.0), Inches(1.8), Inches(6), Inches(4.2),
               ["フレーズ", "件数"], rows, header_fill=GREEN)
    else:
        _textbox(slide, Inches(7.0), Inches(1.9), Inches(6), Inches(1.0),
                 "（口コミ件数が少なく、頻出フレーズを抽出できませんでした）",
                 size=13, color=GREY)


def _slide_insights_sw(prs, insights):
    slide = _blank(prs)
    _title_bar(slide, "インサイト①　強み・弱み", "LLMによる分析")

    _textbox(slide, Inches(0.5), Inches(1.3), Inches(6), Inches(0.4),
             "💪 強み", size=18, bold=True, color=GREEN)
    _bullets(slide, Inches(0.5), Inches(1.85), Inches(6.0), Inches(4.8),
             insights.strengths or ["（なし）"], size=15)

    _textbox(slide, Inches(7.0), Inches(1.3), Inches(6), Inches(0.4),
             "⚠️ 弱み", size=18, bold=True, color=RED)
    _bullets(slide, Inches(7.0), Inches(1.85), Inches(6.0), Inches(4.8),
             insights.weaknesses or ["（なし）"], size=15)


def _slide_insights_action(prs, insights):
    slide = _blank(prs)
    _title_bar(slide, "インサイト②　示唆・改善提案", "LLMによる分析")

    _textbox(slide, Inches(0.5), Inches(1.3), Inches(6), Inches(0.4),
             "💡 示唆", size=18, bold=True, color=BLUE)
    _bullets(slide, Inches(0.5), Inches(1.85), Inches(6.0), Inches(4.8),
             insights.implications or ["（なし）"], size=15)

    _textbox(slide, Inches(7.0), Inches(1.3), Inches(6), Inches(0.4),
             "🔧 改善提案", size=18, bold=True, color=ORANGE)
    _bullets(slide, Inches(7.0), Inches(1.85), Inches(6.0), Inches(4.8),
             insights.improvements or ["（なし）"], size=15)


def _slide_topic_score(prs, ts: "topic_score.TopicScoreResult") -> None:
    """独自指標（感情・トピック統合スコア）スライド。"""
    slide = _blank(prs)
    _title_bar(
        slide, "感情・トピック統合スコア",
        "文単位の感情値 × トピック確率を集計した独自指標（N=1 対応）",
    )

    # 総合感情スコア（重み付き）
    ov = ts.weighted_sentiment_100
    _textbox(slide, Inches(0.5), Inches(1.25), Inches(4.2), Inches(0.4),
             "総合感情スコア", size=14, bold=True, color=GREY)
    _textbox(slide, Inches(0.5), Inches(1.6), Inches(4.2), Inches(1.0),
             f"{ov:.0f} / 100", size=40, bold=True, color=MAGENTA)
    _textbox(slide, Inches(0.5), Inches(2.7), Inches(4.2), Inches(0.5),
             f"分析文数 {ts.n_sentences}／レビュー {ts.n_reviews} 件", size=12, color=GREY)
    _textbox(slide, Inches(0.5), Inches(3.15), Inches(4.3), Inches(2.6),
             "各文をポジ／中立／ネガに評価し、トピック確率で重み付けして"
             "トピックごとの評価スコアへ集計しています（50=中立）。",
             size=12, color=NAVY)

    # トピック別テーブル（感情降順）
    topics = ts.sorted_by_sentiment(reverse=True)
    rows = [
        [t.name, f"{t.sentiment_100:.0f}", f"{t.salience_pct:.0f}%", f"{t.weight*100:.0f}%"]
        for t in topics
    ]
    row_h = min(Inches(0.5), Inches(5.4 / max(len(rows), 1)))
    _table(
        slide, Inches(5.1), Inches(1.3), Inches(7.7), row_h * (len(rows) + 1),
        ["トピック（指標軸）", "感情スコア", "言及度", "重み"], rows,
        header_fill=MAGENTA,
        accent={"col": 1, "color": MAGENTA},
    )

    _textbox(slide, Inches(5.1), Inches(1.3) + row_h * (len(rows) + 1) + Inches(0.1),
             Inches(7.7), Inches(0.4),
             "感情スコア: 60以上=ポジ / 40以下=ネガ。言及度: そのトピックが語られた割合。",
             size=10.5, color=GREY)


def _slide_topics(prs, topic_list: list[Topic]) -> None:
    slide = _blank(prs)
    _title_bar(
        slide, "トピック分析",
        f"{len(topic_list)} トピックを抽出（TF-IDF + KMeans / 仮 BERTopic）",
    )

    rows = []
    for t in topic_list[:8]:
        rows.append([
            f"#{t.id + 1}",
            t.label,
            f"{t.count} 件",
            f"{t.share}%",
            "　/　".join(t.keywords[:5]),
        ])
    row_h = min(Inches(0.44), Inches(4.5 / max(len(rows), 1)))
    if rows:
        _table(
            slide, Inches(0.5), Inches(1.4), Inches(12.3), row_h * len(rows),
            ["#", "トピック", "件数", "割合", "代表キーワード"], rows,
            header_fill=BLUE,
        )

    # sample quotes for top-2 topics
    y = Inches(1.4) + row_h * len(rows) + Inches(0.25)
    for t in topic_list[:2]:
        if not t.samples or y > Inches(6.8):
            break
        _textbox(slide, Inches(0.5), y, Inches(12.3), Inches(0.3),
                 f"【{t.label}】", size=12, bold=True, color=NAVY)
        y += Inches(0.32)
        sample = t.samples[0]
        display = (sample[:150] + "…") if len(sample) > 150 else sample
        _textbox(slide, Inches(0.8), y, Inches(11.8), Inches(0.45),
                 f"「{display}」", size=11, color=GREY)
        y += Inches(0.52)


# --------------------------------------------------------------------------- #
# public entry point
# --------------------------------------------------------------------------- #
def _photo_placeholder(slide):
    box = slide.shapes.add_shape(1, Inches(0.5), Inches(1.5), Inches(5.3), Inches(5.0))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = RGBColor(0xE4, 0xE3, 0xDD)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "施設写真"
    run.font.size = Pt(14)
    run.font.color.rgb = GREY


def _slide_profile(prs, conn, target_name, info, photo_bytes):
    """PROFILE スライド：写真＋基本情報テーブル。"""
    info = info or {}
    slide = _blank(prs)
    _title_bar(slide, "分析施設情報", "PROFILE")

    frow = conn.execute(
        "SELECT id, category FROM facility WHERE name = ?", (target_name,)
    ).fetchone()
    fid = frow["id"] if frow else None
    category = (frow["category"] if frow and frow["category"] else info.get("category")) or "—"
    n_reviews = conn.execute(
        "SELECT COUNT(*) FROM review WHERE facility_id = ?", (fid,)
    ).fetchone()[0] if fid else 0
    avg = conn.execute(
        "SELECT AVG(rating) FROM review WHERE facility_id = ? AND rating IS NOT NULL", (fid,)
    ).fetchone()[0] if fid else None
    review_line = (f"平均 ★{round(avg, 1)} ／ {n_reviews:,} 件"
                   if avg is not None else f"{n_reviews:,} 件")

    # photo (left)
    if photo_bytes:
        try:
            slide.shapes.add_picture(io.BytesIO(photo_bytes), Inches(0.5), Inches(1.5),
                                     width=Inches(5.3))
        except Exception:
            _photo_placeholder(slide)
    else:
        _photo_placeholder(slide)

    # basic info table (right)
    rows = [
        ["施設名", target_name],
        ["業種", category],
        ["住所", info.get("address") or "—"],
        ["アクセス", info.get("access") or "—"],
        ["開業", info.get("open_year") or "—"],
        ["延床", info.get("floor_area") or "—"],
        ["口コミ", review_line],
    ]
    _table(slide, Inches(6.2), Inches(1.5), Inches(6.6), Inches(0.75 * len(rows)),
           ["項目", "内容"], rows, header_fill=MAGENTA, accent={"col": 0, "color": NAVY})


def _slide_symbolic(prs, ranked: list[dict]):
    """象徴的な口コミランキング（TF-IDF総合スコア）。"""
    slide = _blank(prs)
    _title_bar(slide, "象徴的な口コミ ランキング",
               "施設の特徴語をどれだけ体現しているかで口コミを総合スコア化")
    rows = []
    for r in ranked:
        snippet = (r["text"][:64] + "…") if len(r["text"]) > 64 else r["text"]
        star = f'★{r["rating"]}' if r.get("rating") is not None else "-"
        rows.append([f'#{r["rank"]}', f'{r["share"]:.0f}', star, snippet,
                     "・".join(r.get("keywords", [])[:4])])
    if rows:
        _table(slide, Inches(0.4), Inches(1.4), Inches(12.5),
               Inches(min(0.9, 4.8 / len(rows)) * len(rows)),
               ["順位", "象徴度", "評価", "口コミ", "特徴語"], rows,
               header_fill=MAGENTA, accent={"col": 1, "color": MAGENTA})


def _facility_column(slide, left, top, width, height, names):
    """1カラム分の施設名（マゼンタ●＋濃色テキスト）。"""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, nm in enumerate(names):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r1 = p.add_run()
        r1.text = "● "
        r1.font.size = Pt(12)
        r1.font.bold = True
        r1.font.color.rgb = MAGENTA
        r2 = p.add_run()
        r2.text = nm
        r2.font.size = Pt(14)
        r2.font.color.rgb = NAVY
    return box


def _slide_disclaimer(prs):
    """免責事項 — レポート冒頭（1枚目）。"""
    slide = _blank(prs)
    _title_bar(slide, config.DISCLAIMER_TITLE,
               "免責事項 — ご覧いただく前にご確認ください")
    items = [f"{i}. {p}" for i, p in enumerate(config.DISCLAIMER_POINTS, 1)]
    _bullets(slide, Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.4),
             items, size=14, marker="", color=NAVY, space_after=12)
    _textbox(slide, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.4),
             "本レポートは参考情報です（VoiceBAUM）", size=11, color=GREY)


def _slide_appendix(prs, conn, target_name):
    """APPENDIX: 比較対象（ピア）施設の一覧。ピアが無ければスライドを作らない。"""
    try:
        # 口コミが1件以上ある比較施設のみ（プレビューのピア集合に合わせる）
        rows = conn.execute(
            "SELECT f.name, COUNT(r.id) c FROM facility f "
            "LEFT JOIN review r ON r.facility_id = f.id "
            "WHERE f.type = 'comparison' AND f.name <> ? "
            "GROUP BY f.id HAVING c > 0 ORDER BY f.name",
            (target_name,),
        ).fetchall()
        peers = [r[0] for r in rows]
    except Exception:
        peers = []
    if not peers:
        return
    slide = _blank(prs)
    _title_bar(slide, "比較対象施設一覧",
               "同市場の施設からも口コミを抽出し、比較して特徴点を可視化")
    # 施設数（ヘッダー帯の右端に白文字）
    _textbox(slide, Inches(9.9), Inches(0.0), Inches(3.2), Inches(1.0),
             f"{len(peers)} 施設", size=18, bold=True, color=WHITE,
             align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    CAP = 54                       # 3カラム×約18行に収める
    overflow = max(0, len(peers) - CAP)
    shown = peers[:CAP]
    per = (len(shown) + 2) // 3
    cols = [shown[0:per], shown[per:2 * per], shown[2 * per:]]
    lefts = [Inches(0.5), Inches(4.77), Inches(9.04)]
    for chunk, left in zip(cols, lefts):
        if chunk:
            _facility_column(slide, left, Inches(1.5), Inches(4.1), Inches(5.4), chunk)
    if overflow:
        _textbox(slide, Inches(0.5), Inches(7.0), Inches(6), Inches(0.4),
                 f"ほか {overflow} 施設", size=12, bold=True, color=MAGENTA)


def build_report(
    conn: sqlite3.Connection,
    target_name: str,
    axis: str = "comparison_avg",
    specific_name: Optional[str] = None,
    insights: Optional[InsightResult] = None,
    topic_list: Optional[list[Topic]] = None,
    topic_score_result: "Optional[topic_score.TopicScoreResult]" = None,
    profile_info: Optional[dict] = None,
    photo_bytes: Optional[bytes] = None,
    peers: Optional[list[str]] = None,
    output_path: str | Path = "report.pptx",
) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    comp = analysis.build_comparison(
        conn, target_name, axis, specific_name=specific_name, peers=peers
    )
    # 指定競合モード（peers あり）では全施設平均へのフォールバックをしない。
    # 選択した競合にスコアが無いのに「全体平均」を出すと比較軸がすり替わるため。
    if comp is None and not peers:
        comp = analysis.build_comparison(conn, target_name, "all_avg")
    profile = text_analysis.build_profile(conn, target_name, top_n=20)

    # 独自指標（感情・トピック統合スコア）。未算出なら内部で算出。
    if topic_score_result is None:
        try:
            topic_score_result = topic_score.analyze_facility(conn, target_name)
        except Exception:
            topic_score_result = None

    if insights is None:
        insights = InsightResult(summary="（インサイト未生成。⑤画面で生成すると反映されます）")
    # stash review count for the summary cards
    insights._n_reviews = f"{profile.n_reviews} 件"

    baseline_label = comp.baseline_label if comp else "（比較データなし）"

    _slide_disclaimer(prs)                       # 冒頭1枚目
    _slide_title(prs, target_name, baseline_label)
    _slide_summary(prs, target_name, insights, comp)
    _slide_profile(prs, conn, target_name, profile_info, photo_bytes)
    if topic_score_result is not None and not topic_score_result.empty:
        _slide_topic_score(prs, topic_score_result)
    if comp is not None:
        _slide_score_charts(prs, comp)
        _slide_top5(prs, comp)
    if not profile.empty:
        _slide_text_analysis(prs, profile)
        try:
            _rr = conn.execute(
                "SELECT rating, text FROM review WHERE facility_id = "
                "(SELECT id FROM facility WHERE name = ?) "
                "AND text IS NOT NULL AND text != ''", (target_name,)
            ).fetchall()
            _ranked = text_analysis.symbolic_ranking(
                [(x[0], x[1]) for x in _rr], profile.tfidf_keywords, top_k=5
            )
        except Exception:
            _ranked = []
        if _ranked:
            _slide_symbolic(prs, _ranked)
    if topic_list:
        _slide_topics(prs, topic_list)
    _slide_insights_sw(prs, insights)
    _slide_insights_action(prs, insights)
    _slide_appendix(prs, conn, target_name)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path
