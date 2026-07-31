"""Tests for the SLIDE 1「施設・基本情報」統合スライド（preview.html_facility_info）と、
build_bundle が新たに生成する floor_area / review_trend / peer_display。
"""
from pathlib import Path

from pptx import Presentation

from src import db, preview, report, topic_score
from src.review_csv import ParsedReview

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


# --------------------------------------------------------------------------- #
# preview.html_facility_info（純関数）
# --------------------------------------------------------------------------- #
def _min_bundle(**overrides):
    b = {
        "target": "テスト施設", "category": "美術館・博物館", "n_reviews": 120,
        "avg_rating": 4.2, "address": "東京都○○区1-1-1", "open_year": "2015年",
        "floor_area": "28,500㎡", "photo_data_uri": None,
        "review_trend": [("2025-01", 10), ("2025-02", 25), ("2025-03", 50)],
        "peer_display": [{"name": "比較館A", "photo_data_uri": None},
                          {"name": "比較館B", "photo_data_uri": None}],
        "peer_display_same_category": True,
    }
    b.update(overrides)
    return b


def test_facility_info_contains_core_fields():
    html = preview.html_facility_info(_min_bundle())
    assert "施設・基本情報" in html
    assert "施設プロフィール" in html
    assert "口コミサマリー" in html
    assert "口コミ数の推移" in html
    assert "比較対象施設" in html
    assert "東京都○○区1-1-1" in html
    assert "2015年" in html
    assert "28,500㎡" in html
    assert "美術館・博物館" in html
    assert "120" in html          # n_reviews
    assert "4.2" in html          # avg_rating
    assert "比較館A" in html and "比較館B" in html


def test_facility_info_escapes_html():
    html = preview.html_facility_info(_min_bundle(
        category="<b>x</b>&y",
        peer_display=[{"name": "<script>alert(1)</script>", "photo_data_uri": None}],
    ))
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html
    assert "<b>x</b>&y" not in html


def test_facility_info_missing_fields_show_placeholder():
    html = preview.html_facility_info(_min_bundle(
        address=None, open_year=None, floor_area=None, category="—",
    ))
    assert "施設写真" in html            # 写真プレースホルダ
    assert "—" in html                   # 各項目のダッシュ表示


def test_facility_info_empty_trend_shows_fallback():
    html = preview.html_facility_info(_min_bundle(review_trend=[]))
    assert "データが不足しています" in html


def test_facility_info_single_point_trend_shows_fallback():
    html = preview.html_facility_info(_min_bundle(review_trend=[("2025-01", 5)]))
    assert "データが不足しています" in html


def test_facility_info_empty_peers_shows_fallback():
    html = preview.html_facility_info(_min_bundle(peer_display=[]))
    assert "比較対象施設がありません" in html


def test_facility_info_peer_note_reflects_same_category_flag():
    html_same = preview.html_facility_info(_min_bundle(peer_display_same_category=True))
    assert "同カテゴリの2施設を比較対象として設定" in html_same

    html_other = preview.html_facility_info(_min_bundle(peer_display_same_category=False))
    assert "※2施設を比較対象として設定" in html_other
    assert "同カテゴリ" not in html_other.split("比較対象施設（同カテゴリの類似施設）")[-1].split("施設を比較対象")[0]


# --------------------------------------------------------------------------- #
# 増減傾向の簡易判定（_trend_trend_note）
# --------------------------------------------------------------------------- #
def test_trend_note_detects_increase():
    # 累積: 前半3ヶ月ほぼ横ばい → 後半3ヶ月で急増
    trend = [("m1", 10), ("m2", 20), ("m3", 30), ("m4", 60), ("m5", 100), ("m6", 150), ("m7", 210)]
    assert preview._trend_trend_note(trend) == "継続的に増加傾向"


def test_trend_note_detects_decrease():
    trend = [("m1", 10), ("m2", 60), ("m3", 120), ("m4", 130), ("m5", 138), ("m6", 144), ("m7", 148)]
    assert preview._trend_trend_note(trend) == "減少傾向"


def test_trend_note_too_short_returns_empty():
    assert preview._trend_trend_note([("m1", 5), ("m2", 10)]) == ""


# --------------------------------------------------------------------------- #
# build_bundle が floor_area / review_trend / peer_display を正しく組み立てるか
# --------------------------------------------------------------------------- #
def _seed(tmp_path):
    conn = db.get_conn(tmp_path / "t.db")
    db.init_db(conn)
    db.upsert_facility(conn, "対象館", ftype="target", category="企業ミュージアム",
                       floor_area="12,000㎡")
    for i, m in enumerate(["2025-01", "2025-02", "2025-03"]):
        db.insert_reviews(conn, db.upsert_facility(conn, "対象館"), [
            ParsedReview(review_id=f"t{i}_{j}", rating=5, text="良い",
                         review_date=f"{m}-01T00:00:00Z", reviewer_name="A",
                         local_guide=False, likes=None, owner_response="",
                         owner_response_date="", subscores=[])
            for j in range(3)
        ])
    for nm, cat in [("同カテゴリ館1", "企業ミュージアム"), ("同カテゴリ館2", "企業ミュージアム"),
                    ("別カテゴリ館", "水族館")]:
        fid = db.upsert_facility(conn, nm, ftype="comparison", category=cat)
        db.insert_reviews(conn, fid, [
            ParsedReview(review_id=f"{nm}_1", rating=4, text="普通",
                         review_date="2025-02-01T00:00:00Z", reviewer_name="B",
                         local_guide=False, likes=None, owner_response="",
                         owner_response_date="", subscores=[])
        ])
    return conn


def test_build_bundle_includes_floor_area_and_trend(tmp_path):
    conn = _seed(tmp_path)
    names = ["対象館", "同カテゴリ館1", "同カテゴリ館2", "別カテゴリ館"]
    topic_results = {n: topic_score.analyze_facility(conn, n) for n in names}
    bundle = preview.build_bundle(conn, "対象館", topic_results, None, None)

    assert bundle["floor_area"] == "12,000㎡"
    assert bundle["review_trend"] == [("2025-01", 3), ("2025-02", 6), ("2025-03", 9)]  # 累積


def test_build_bundle_prefers_same_category_peers(tmp_path):
    conn = _seed(tmp_path)
    names = ["対象館", "同カテゴリ館1", "同カテゴリ館2", "別カテゴリ館"]
    topic_results = {n: topic_score.analyze_facility(conn, n) for n in names}
    bundle = preview.build_bundle(conn, "対象館", topic_results, None, None)

    peer_names = [p["name"] for p in bundle["peer_display"]]
    assert "同カテゴリ館1" in peer_names and "同カテゴリ館2" in peer_names
    assert bundle["peer_display_same_category"] is True
    # 同カテゴリが先頭に来る（別カテゴリは補充要員）
    assert peer_names.index("同カテゴリ館1") < peer_names.index("別カテゴリ館")


# --------------------------------------------------------------------------- #
# report._slide_profile（PPTX）— 延床行の反映
# --------------------------------------------------------------------------- #
def _all_text(prs) -> str:
    out = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if sh.has_table:
                for row in sh.table.rows:
                    out.append(" ".join(c.text for c in row.cells))
    return "\n".join(out)


def test_pptx_profile_includes_floor_area(tmp_path):
    conn = db.get_conn(tmp_path / "t.db")
    db.init_db(conn)
    fid = db.upsert_facility(conn, "館A", ftype="target")
    db.insert_reviews(conn, fid, [
        ParsedReview(review_id="r1", rating=5, text="良い展示でした。",
                     review_date="2025-01-01", reviewer_name="u", local_guide=False,
                     likes=0, owner_response="", owner_response_date="", subscores=[]),
    ])
    out = report.build_report(
        conn, "館A", profile_info={"floor_area": "9,800㎡", "address": "東京都"},
        output_path=tmp_path / "r.pptx",
    )
    prs = Presentation(str(out))
    blob = _all_text(prs)
    assert "延床" in blob
    assert "9,800㎡" in blob


def test_pptx_profile_floor_area_placeholder_when_missing(tmp_path):
    conn = db.get_conn(tmp_path / "t2.db")
    db.init_db(conn)
    fid = db.upsert_facility(conn, "館B", ftype="target")
    db.insert_reviews(conn, fid, [
        ParsedReview(review_id="r1", rating=4, text="普通でした。",
                     review_date="2025-01-01", reviewer_name="u", local_guide=False,
                     likes=0, owner_response="", owner_response_date="", subscores=[]),
    ])
    out = report.build_report(conn, "館B", output_path=tmp_path / "r.pptx")
    prs = Presentation(str(out))
    blob = _all_text(prs)
    assert "延床" in blob   # ラベル行自体は常に出る（値は「—」）
