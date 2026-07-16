"""Tests for the disclaimer slide (preview HTML + PPTX first slide)."""
import tempfile
from pathlib import Path

from pptx import Presentation

from src import config, db, preview, report
from src.review_csv import ParsedReview


def test_html_disclaimer_contains_all_points():
    html = preview.html_disclaimer({"date": "2026年07月16日"})
    assert config.DISCLAIMER_TITLE in html
    assert "免責事項" in html
    for p in config.DISCLAIMER_POINTS:
        # 先頭20文字が含まれる（エスケープ差異を避けて部分一致）
        assert p[:20] in html
    assert "2026年07月16日" in html


def test_pptx_disclaimer_is_first_slide():
    tmp = Path(tempfile.mkdtemp())
    conn = db.get_conn(tmp / "t.db"); db.init_db(conn)
    fid = db.upsert_facility(conn, "対象館", ftype="target")
    db.insert_reviews(conn, fid, [
        ParsedReview(f"r{i}", 5, "良い展示だった。", f"2025-01-0{i+1}",
                     "u", False, 0, "", "", []) for i in range(5)])
    out = report.build_report(conn, "対象館", output_path=tmp / "r.pptx")
    prs = Presentation(str(out))
    first = "\n".join(sh.text_frame.text for sh in prs.slides[0].shapes
                      if sh.has_text_frame)
    assert config.DISCLAIMER_TITLE in first
    assert "免責事項" in first
