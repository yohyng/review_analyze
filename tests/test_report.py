"""Tests for step 9: PPTX report generation."""
from pathlib import Path

from pptx import Presentation

from src import db, report, score_excel
from src.llm import InsightResult
from src.review_csv import ParsedReview

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def _seed(tmp_path):
    conn = db.get_conn(tmp_path / "t.db")
    db.init_db(conn)

    df = score_excel.read_table(SAMPLES / "sample_scores.xlsx")
    g = score_excel.guess_columns(df)
    scores = score_excel.extract_scores_wide(df, g.name_col, g.numeric_cols)

    fid_t = db.upsert_facility(conn, "風の海", ftype="target")
    db.upsert_scores(conn, fid_t, scores["風の海"], scale=5)
    db.insert_reviews(conn, fid_t, [
        ParsedReview("r1", 5, "スタッフが丁寧で景色も最高でした。", "2025-01-01",
                     "u", True, 0, "", "", []),
        ParsedReview("r2", 3, "設備が古く食事も遅かった。", "2025-02-01",
                     "u2", False, 0, "", "", []),
    ])

    fid_c = db.upsert_facility(conn, "海の宿 潮", ftype="comparison")
    db.upsert_scores(conn, fid_c, scores["海の宿 潮"], scale=5)
    return conn


def test_build_report_creates_valid_pptx(tmp_path):
    conn = _seed(tmp_path)
    out = report.build_report(conn, "風の海", output_path=tmp_path / "r.pptx")
    assert out.exists()

    prs = Presentation(str(out))
    # title + summary + score charts + top5 + text + 2 insight slides
    assert len(prs.slides) >= 6


def test_report_contains_charts_and_tables(tmp_path):
    conn = _seed(tmp_path)
    out = report.build_report(conn, "風の海", output_path=tmp_path / "r.pptx")
    prs = Presentation(str(out))

    n_charts = sum(sh.has_chart for s in prs.slides for sh in s.shapes)
    n_tables = sum(sh.has_table for s in prs.slides for sh in s.shapes)
    assert n_charts >= 2          # radar + bar
    assert n_tables >= 2          # strengths + weaknesses


def test_report_with_insights(tmp_path):
    conn = _seed(tmp_path)
    insights = InsightResult(
        summary="テスト要約",
        strengths=["強み1", "強み2"],
        weaknesses=["弱み1"],
        implications=["示唆1"],
        improvements=["改善1", "改善2"],
    )
    out = report.build_report(conn, "風の海", insights=insights,
                              output_path=tmp_path / "r.pptx")
    assert out.exists()


def test_report_without_scores_text_only(tmp_path):
    """A facility with only reviews (no Excel scores) still produces a deck."""
    conn = db.get_conn(tmp_path / "t.db")
    db.init_db(conn)
    fid = db.upsert_facility(conn, "テキストのみ施設", ftype="target")
    db.insert_reviews(conn, fid, [
        ParsedReview("x1", 4, "立地が良くスタッフも親切でした。", "2025-01-01",
                     "u", False, 0, "", "", []),
    ])
    out = report.build_report(conn, "テキストのみ施設", output_path=tmp_path / "r.pptx")
    assert out.exists()
    prs = Presentation(str(out))
    assert len(prs.slides) >= 3   # title + summary + insights at minimum


# --------------------------------------------------------------------------- #
# 設定画面の検索窓
#
#   <input> に直接 height を当てると、Streamlit 側の器（stTextInput）は既定の
#   高さのままなので、はみ出したぶんが次の要素に隠れて下端が切れる。
#   見た目は BaseWeb のラッパ（stTextInputRootElement）側に持たせること。
#   実ブラウザで測って直したときの取り決めを、ここで固定する。
# --------------------------------------------------------------------------- #
def _setup_css() -> str:
    from pathlib import Path

    from src.ui import analysis_mode

    src = Path(analysis_mode.__file__).read_text(encoding="utf-8")
    start = src.index("# Hero search field:")
    # 生HTMLは _html()（= st.html）で流す。Markdown を通さないため。
    # v0.56.0 以前は st.markdown(..., unsafe_allow_html=True) だった。
    return src[start:src.index('""")', start)]


def test_search_box_styles_the_wrapper_not_the_input():
    css = _setup_css()
    assert 'div[data-baseweb="input"]{' in css, (
        "枠・角丸・影はラッパ側に当てること（input に直接だと器からはみ出す）"
    )
    tail = css.split('div[data-testid="stTextInput"] input{')[-1]
    assert "height:60px" not in tail, "input 側に固定の高さを戻さない"


def test_search_box_container_reserves_room_for_the_frame():
    """器の高さを確保しないと、枠の下端が次の要素に隠れる。"""
    css = _setup_css()
    assert 'div[data-testid="stTextInput"]{ min-height:64px!important; }' in css


def test_inner_base_input_is_transparent():
    """内側の base-input は白で塗るので、透過させないと角丸と虫めがねが隠れる。"""
    css = _setup_css()
    assert 'div[data-baseweb="base-input"]{' in css
    assert "background:transparent!important" in css
