"""Tests for the APPENDIX 比較対象施設一覧 slide (preview HTML + PPTX)."""
from pathlib import Path

from pptx import Presentation

from src import db, preview, report
from src.review_csv import ParsedReview


# --------------------------------------------------------------------------- #
# preview.html_appendix（純関数）
# --------------------------------------------------------------------------- #
def _bundle(names):
    return {"peer_names": list(names), "peer_count": len(names)}


def test_appendix_empty_returns_blank():
    assert preview.html_appendix({"peer_names": [], "peer_count": 0}) == ""
    assert preview.html_appendix({}) == ""


def test_appendix_lists_all_names_and_count():
    names = ["施設A", "施設B", "施設C"]
    html = preview.html_appendix(_bundle(names))
    assert "比較対象施設一覧" in html
    assert "APPENDIX" in html
    assert "3 施設" in html
    for n in names:
        assert n in html


def test_appendix_escapes_html():
    html = preview.html_appendix(_bundle(["<b>x</b>&y"]))
    assert "<b>x</b>" not in html          # 生タグは出さない
    assert "&lt;b&gt;" in html


def test_appendix_overflow_caps_and_notes():
    names = [f"施設{i:02d}" for i in range(60)]
    html = preview.html_appendix(_bundle(names))
    assert "60 施設" in html               # 総数は正しく表示
    assert "ほか 15 施設" in html           # 60 - 上限45
    assert "施設00" in html and "施設44" in html
    assert "施設45" not in html            # 上限で切られている


# --------------------------------------------------------------------------- #
# report._slide_appendix（PPTX）
# --------------------------------------------------------------------------- #
def _seed_with_peers(tmp_path):
    conn = db.get_conn(tmp_path / "t.db")
    db.init_db(conn)
    ft = db.upsert_facility(conn, "対象館", ftype="target")
    db.insert_reviews(conn, ft, [
        ParsedReview("t1", 5, "とても良い展示でした。", "2025-01-01", "u", False, 0, "", "", []),
        ParsedReview("t2", 4, "スタッフが親切で気持ちよかった。", "2025-01-02", "u", False, 0, "", "", []),
    ])
    for i, nm in enumerate(["比較館A", "比較館B", "比較館C"]):
        fc = db.upsert_facility(conn, nm, ftype="comparison")
        db.insert_reviews(conn, fc, [
            ParsedReview(f"c{i}", 4, "普通の展示だった。", "2025-01-01", "u", False, 0, "", "", []),
        ])
    # レビュー0件の比較施設 — APPENDIX には出ないはず
    db.upsert_facility(conn, "口コミ無し館", ftype="comparison")
    return conn


def _all_text(prs) -> str:
    out = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
    return "\n".join(out)


def test_pptx_appendix_slide_present_and_last(tmp_path):
    conn = _seed_with_peers(tmp_path)
    out = report.build_report(conn, "対象館", output_path=tmp_path / "r.pptx")
    prs = Presentation(str(out))

    blob = _all_text(prs)
    assert "比較対象施設一覧" in blob
    assert "比較館A" in blob and "比較館B" in blob and "比較館C" in blob
    assert "3 施設" in blob
    assert "口コミ無し館" not in blob        # レビュー0件は除外

    # 「最後に」— 末尾スライドに置かれている
    last = prs.slides[-1]
    last_text = "\n".join(sh.text_frame.text for sh in last.shapes if sh.has_text_frame)
    assert "比較対象施設一覧" in last_text


def test_pptx_no_appendix_when_no_peers(tmp_path):
    conn = db.get_conn(tmp_path / "t2.db")
    db.init_db(conn)
    ft = db.upsert_facility(conn, "単体館", ftype="target")
    db.insert_reviews(conn, ft, [
        ParsedReview("s1", 5, "良い。", "2025-01-01", "u", False, 0, "", "", []),
    ])
    out = report.build_report(conn, "単体館", output_path=tmp_path / "r.pptx")
    prs = Presentation(str(out))
    assert "比較対象施設一覧" not in _all_text(prs)
